"""Safe, schema-mapped student imports from CSV and Excel workbooks."""

from __future__ import annotations

import csv
import io
import re
import unicodedata
import zipfile
from decimal import Decimal, InvalidOperation
from pathlib import Path
from urllib.parse import urlparse

import openpyxl
import xlrd
from django.core.exceptions import ValidationError as DjangoValidationError
from django.core.validators import validate_email
from django.db.models.functions import Lower

from apps.users.models import User
from .models import StudentProfile


MAX_IMPORT_FILE_SIZE = 5 * 1024 * 1024
MAX_IMPORT_ROWS = 2_000
MAX_IMPORT_COLUMNS = 100
MAX_UNCOMPRESSED_XLSX_SIZE = 50 * 1024 * 1024
USER_FIRST_NAME_MAX_LENGTH = User._meta.get_field('first_name').max_length
USER_LAST_NAME_MAX_LENGTH = User._meta.get_field('last_name').max_length
USER_EMAIL_MAX_LENGTH = User._meta.get_field('email').max_length
PORTFOLIO_URL_MAX_LENGTH = StudentProfile._meta.get_field('portfolio_google_docs_url').max_length


class StudentImportError(ValueError):
    pass


IMPORT_FIELDS = (
    {'key': 'full_name', 'label': 'Full name', 'required': True},
    {'key': 'grade', 'label': 'Grade', 'required': True},
    {'key': 'email', 'label': 'Email', 'required': False},
    {'key': 'phone', 'label': 'Student phone', 'required': False},
    {'key': 'parent_name', 'label': 'Parent name', 'required': False},
    {'key': 'parent_phone', 'label': 'Parent phone', 'required': False},
    {'key': 'gpa', 'label': 'GPA', 'required': False},
    {'key': 'ielts_score', 'label': 'IELTS', 'required': False},
    {'key': 'sat_score', 'label': 'SAT', 'required': False},
    {'key': 'target_major', 'label': 'Target major', 'required': False},
    {'key': 'target_countries', 'label': 'Target countries', 'required': False},
    {'key': 'budget_usd', 'label': 'Annual budget USD', 'required': False},
    {'key': 'scholarship_needed', 'label': 'Scholarship needed', 'required': False},
    {'key': 'status', 'label': 'Account status', 'required': False},
    {'key': 'portfolio_google_docs_url', 'label': 'Google Docs portfolio', 'required': False},
)


def _compact(value):
    text = unicodedata.normalize('NFKD', str(value or '').casefold())
    text = ''.join(char for char in text if not unicodedata.combining(char))
    text = text.replace('ʻ', '').replace('ʼ', '').replace('’', '').replace("'", '')
    return re.sub(r'[^a-zа-я0-9]+', '', text)


HEADER_ALIASES = {
    'full_name': {
        'fish', 'fio', 'fullname', 'studentname', 'oquvchifish', 'oquvchifio',
        'familiyaismsharif', 'familiyaimyaotchestvo',
    },
    'grade': {'sinf', 'grade', 'class', 'klass'},
    'email': {'email', 'emailaddress', 'elektronpochta', 'pochta'},
    # Generic "Telefon raqami" is intentionally not guessed: contract sheets
    # often contain several parent/student phone columns with the same label.
    'phone': {'studentphone', 'studenttelefon', 'oquvchitelefoni', 'oquvchitelefonraqami'},
    'parent_name': {
        'parentname', 'parentfish', 'otaonafish', 'otayokionaningfish',
        'otayokionasiningfish', 'roditelfio',
    },
    'parent_phone': {
        'parentphone', 'otaonatelefoni', 'otayokionaningtelefonraqami',
        'otayokionasiningtelefonraqami', 'telefonroditelya',
    },
    'gpa': {'gpa', 'averagegrade', 'ortachabaho'},
    'ielts_score': {'ielts', 'ieltsscore', 'ieltsball'},
    'sat_score': {'sat', 'satscore', 'satball'},
    'target_major': {'targetmajor', 'major', 'yonalish', 'mutaxassislik'},
    'target_countries': {'targetcountries', 'countries', 'davlatlar', 'mamlakatlar'},
    'budget_usd': {'budgetusd', 'annualbudgetusd', 'budget', 'yillikbudjet'},
    'scholarship_needed': {'scholarshipneeded', 'grantkerak', 'stipendiyakerak'},
    'status': {'holati', 'status', 'accountstatus'},
    'portfolio_google_docs_url': {
        'googledocs', 'googledocslink', 'googledocsurl', 'portfoliolink',
        'portfoliogoogleDocsurl'.casefold(),
    },
}


def _cell_text(value):
    if value is None:
        return ''
    if hasattr(value, 'isoformat'):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _header_score(row):
    aliases = set().union(*HEADER_ALIASES.values())
    return sum(1 for value in row if _compact(value) in aliases)


def _choose_table(tables):
    best = None
    for sheet_name, rows in tables:
        candidates = [(index, _header_score(row)) for index, row in enumerate(rows[:20]) if any(_cell_text(cell) for cell in row)]
        if not candidates:
            continue
        header_index, score = max(candidates, key=lambda item: (item[1], -item[0]))
        candidate = (score, sheet_name, header_index, rows)
        if best is None or candidate[0] > best[0]:
            best = candidate
    if best is None:
        raise StudentImportError('The selected file does not contain a readable table.')
    _, sheet_name, header_index, rows = best
    return sheet_name, header_index, rows


def _csv_tables(data):
    try:
        text = data.decode('utf-8-sig')
    except UnicodeDecodeError as exc:
        raise StudentImportError('CSV files must use UTF-8 encoding.') from exc
    try:
        dialect = csv.Sniffer().sniff(text[:8192], delimiters=',;\t')
    except csv.Error:
        dialect = csv.excel
    rows = list(csv.reader(io.StringIO(text), dialect))
    return [('CSV', rows[:MAX_IMPORT_ROWS + 21])]


def _xlsx_tables(data):
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            if len(archive.infolist()) > 2_000 or sum(item.file_size for item in archive.infolist()) > MAX_UNCOMPRESSED_XLSX_SIZE:
                raise StudentImportError('The Excel workbook expands beyond the safe import limit.')
    except zipfile.BadZipFile as exc:
        raise StudentImportError('The Excel workbook is damaged or invalid.') from exc

    try:
        workbook = openpyxl.load_workbook(
            io.BytesIO(data), read_only=True, data_only=True, keep_links=False,
        )
    except (OSError, ValueError, KeyError, zipfile.BadZipFile) as exc:
        raise StudentImportError('The Excel workbook could not be read.') from exc
    tables = []
    try:
        for sheet in workbook.worksheets[:20]:
            if sheet.max_column > MAX_IMPORT_COLUMNS:
                raise StudentImportError(f'Worksheet "{sheet.title}" has more than {MAX_IMPORT_COLUMNS} columns.')
            rows = list(sheet.iter_rows(
                min_row=1,
                max_row=min(sheet.max_row, MAX_IMPORT_ROWS + 21),
                max_col=min(sheet.max_column, MAX_IMPORT_COLUMNS),
                values_only=True,
            ))
            tables.append((sheet.title, rows))
    finally:
        workbook.close()
    return tables


def _xls_tables(data):
    try:
        workbook = xlrd.open_workbook(file_contents=data, on_demand=True)
    except (xlrd.XLRDError, OSError, ValueError) as exc:
        raise StudentImportError('The legacy Excel workbook could not be read.') from exc
    tables = []
    try:
        for sheet in workbook.sheets()[:20]:
            if sheet.ncols > MAX_IMPORT_COLUMNS:
                raise StudentImportError(f'Worksheet "{sheet.name}" has more than {MAX_IMPORT_COLUMNS} columns.')
            rows = [sheet.row_values(index, 0, min(sheet.ncols, MAX_IMPORT_COLUMNS)) for index in range(min(sheet.nrows, MAX_IMPORT_ROWS + 21))]
            tables.append((sheet.name, rows))
    finally:
        workbook.release_resources()
    return tables


def read_import_table(upload):
    if not upload:
        raise StudentImportError('Select an Excel or CSV file.')
    if upload.size <= 0:
        raise StudentImportError('The selected file is empty.')
    if upload.size > MAX_IMPORT_FILE_SIZE:
        raise StudentImportError('The import file is larger than the 5 MB limit.')
    extension = Path(upload.name or '').suffix.lower()
    if extension not in {'.xlsx', '.xls', '.csv'}:
        raise StudentImportError('Use an .xlsx, .xls, or .csv file.')
    data = upload.read(MAX_IMPORT_FILE_SIZE + 1)
    upload.seek(0)
    if len(data) > MAX_IMPORT_FILE_SIZE:
        raise StudentImportError('The import file is larger than the 5 MB limit.')
    if extension == '.csv':
        tables = _csv_tables(data)
    elif extension == '.xlsx':
        tables = _xlsx_tables(data)
    else:
        tables = _xls_tables(data)
    sheet_name, header_index, raw_rows = _choose_table(tables)
    raw_header = raw_rows[header_index]
    if len(raw_header) > MAX_IMPORT_COLUMNS:
        raise StudentImportError(f'Import files may contain at most {MAX_IMPORT_COLUMNS} columns.')
    headers = [
        {
            'key': f'column_{index}',
            'label': _cell_text(value) or f'Column {index + 1}',
            'index': index,
            'position': index + 1,
        }
        for index, value in enumerate(raw_header)
    ]
    data_rows = []
    truncated = False
    for offset, raw_row in enumerate(raw_rows[header_index + 1:], start=header_index + 2):
        if not any(_cell_text(value) for value in raw_row):
            continue
        if len(data_rows) >= MAX_IMPORT_ROWS:
            truncated = True
            break
        data_rows.append({'row_number': offset, 'cells': list(raw_row[:len(headers)])})
    if truncated:
        raise StudentImportError(f'Import files may contain at most {MAX_IMPORT_ROWS} data rows.')
    if not data_rows:
        raise StudentImportError('No student rows were found below the header row.')
    return {'sheet_name': sheet_name, 'headers': headers, 'rows': data_rows}


def suggested_mapping(headers):
    suggestions = {}
    claimed = set()
    for field in IMPORT_FIELDS:
        key = field['key']
        aliases = HEADER_ALIASES.get(key, set())
        for header in headers:
            if header['key'] in claimed:
                continue
            if _compact(header['label']) in aliases:
                suggestions[key] = header['key']
                claimed.add(header['key'])
                break
    return suggestions


def resolve_mapping(headers, requested):
    available = {header['key'] for header in headers}
    inferred = suggested_mapping(headers)
    requested = requested if isinstance(requested, dict) else {}
    resolved = {}
    for field in IMPORT_FIELDS:
        key = field['key']
        selected = requested[key] if key in requested else inferred.get(key, '')
        resolved[key] = selected if selected in available else ''
    return resolved, inferred


def _mapped_value(raw_row, headers_by_key, mapping, field):
    column_key = mapping.get(field)
    if not column_key:
        return ''
    index = headers_by_key[column_key]['index']
    cells = raw_row['cells']
    return _cell_text(cells[index] if index < len(cells) else '')


def _decimal_value(value, *, label, minimum, maximum, places):
    if value == '':
        return None, None
    normalized = str(value).replace(' ', '').replace(',', '.')
    try:
        number = Decimal(normalized)
    except InvalidOperation:
        return None, f'{label} must be a number.'
    if number < Decimal(str(minimum)) or number > Decimal(str(maximum)):
        return None, f'{label} must be between {minimum} and {maximum}.'
    quantizer = Decimal('1').scaleb(-places)
    return number.quantize(quantizer), None


def _integer_value(value, *, label, minimum, maximum):
    if value == '':
        return None, None
    normalized = str(value).strip().replace(' ', '')
    if re.fullmatch(r'\d{1,3}(,\d{3})+', normalized):
        normalized = normalized.replace(',', '')
    else:
        normalized = normalized.replace(',', '.')
    try:
        number = Decimal(normalized)
        if number != number.to_integral_value():
            raise InvalidOperation
        integer = int(number)
    except (InvalidOperation, ValueError):
        return None, f'{label} must be a whole number.'
    if integer < minimum or integer > maximum:
        return None, f'{label} must be between {minimum} and {maximum}.'
    return integer, None


def _grade_value(value):
    normalized = _compact(value)
    if normalized in {'gap', 'gapyear'}:
        return StudentProfile.Grade.GAP_YEAR, None
    match = re.search(r'(?<!\d)(8|9|10|11)(?!\d)', str(value))
    if match:
        return match.group(1), None
    return '', 'Grade must be 8, 9, 10, 11, or gap year.'


def _boolean_value(value, default=True):
    normalized = _compact(value)
    if not normalized:
        return default
    if normalized in {'0', 'false', 'no', 'yoq', 'нет', 'kerakemas', 'nofaol', 'inactive', 'closed'}:
        return False
    return True


def _valid_google_docs_url(value):
    if not value:
        return True
    parsed = urlparse(value)
    parts = [part for part in parsed.path.split('/') if part]
    return (
        parsed.scheme == 'https'
        and parsed.hostname == 'docs.google.com'
        and len(parts) >= 3
        and parts[0] == 'document'
        and parts[1] == 'd'
        and bool(parts[2])
    )


def _identity_key(value):
    return ' '.join(str(value or '').casefold().split())


def _split_name(full_name):
    parts = full_name.split()
    return parts[0], ' '.join(parts[1:])


def prepare_student_import(upload, *, school, requested_mapping=None, portfolio_links=None):
    table = read_import_table(upload)
    headers_by_key = {header['key']: header for header in table['headers']}
    mapping, inferred = resolve_mapping(table['headers'], requested_mapping)
    portfolio_links = portfolio_links if isinstance(portfolio_links, dict) else {}
    global_errors = [
        f'Select the {field["label"]} column.'
        for field in IMPORT_FIELDS
        if field['required'] and not mapping.get(field['key'])
    ]

    input_emails = set()
    for raw_row in table['rows']:
        email = _mapped_value(raw_row, headers_by_key, mapping, 'email').strip().lower()
        if email:
            input_emails.add(email)
    existing_emails = set(
        User.objects.annotate(email_lower=Lower('email'))
        .filter(email_lower__in=input_emails)
        .values_list('email_lower', flat=True)
    ) if input_emails else set()
    existing_names = {
        _identity_key(' '.join(filter(None, (first_name, last_name))))
        for first_name, last_name in User.objects.filter(
            role=User.Role.STUDENT, school=school,
        ).values_list('first_name', 'last_name')
    }

    seen_emails = set()
    seen_names = set()
    rows = []
    for raw_row in table['rows']:
        row_number = raw_row['row_number']
        errors = []
        warnings = []
        full_name = ' '.join(_mapped_value(raw_row, headers_by_key, mapping, 'full_name').split())
        if not full_name:
            errors.append('Full name is required.')
        else:
            first_name, last_name = _split_name(full_name)
            if len(first_name) > USER_FIRST_NAME_MAX_LENGTH:
                errors.append(f'First name must use {USER_FIRST_NAME_MAX_LENGTH} characters or fewer.')
            if len(last_name) > USER_LAST_NAME_MAX_LENGTH:
                errors.append(f'Last name must use {USER_LAST_NAME_MAX_LENGTH} characters or fewer.')

        grade_raw = _mapped_value(raw_row, headers_by_key, mapping, 'grade')
        grade, grade_error = _grade_value(grade_raw) if grade_raw else ('', 'Grade is required.')
        if grade_error:
            errors.append(grade_error)

        email = _mapped_value(raw_row, headers_by_key, mapping, 'email').strip().lower()
        if email:
            if len(email) > USER_EMAIL_MAX_LENGTH:
                errors.append(f'Email must use {USER_EMAIL_MAX_LENGTH} characters or fewer.')
            try:
                validate_email(email)
            except DjangoValidationError:
                errors.append('Email address is invalid.')
        else:
            warnings.append('An internal login email will be generated.')

        phone = _mapped_value(raw_row, headers_by_key, mapping, 'phone').strip()
        if len(phone) > 32:
            errors.append('Student phone must use 32 characters or fewer.')
        parent_name = _mapped_value(raw_row, headers_by_key, mapping, 'parent_name').strip()
        parent_phone = _mapped_value(raw_row, headers_by_key, mapping, 'parent_phone').strip()
        parent_contact = ' · '.join(value for value in (parent_name, parent_phone) if value)
        if len(parent_contact) > 120:
            errors.append('Parent contact must use 120 characters or fewer.')

        gpa, error = _decimal_value(
            _mapped_value(raw_row, headers_by_key, mapping, 'gpa'),
            label='GPA', minimum=0, maximum=5, places=2,
        )
        if error:
            errors.append(error)
        ielts_score, error = _decimal_value(
            _mapped_value(raw_row, headers_by_key, mapping, 'ielts_score'),
            label='IELTS', minimum=0, maximum=9, places=1,
        )
        if error:
            errors.append(error)
        sat_score, error = _integer_value(
            _mapped_value(raw_row, headers_by_key, mapping, 'sat_score'),
            label='SAT', minimum=400, maximum=1600,
        )
        if error:
            errors.append(error)
        budget_usd, error = _integer_value(
            _mapped_value(raw_row, headers_by_key, mapping, 'budget_usd'),
            label='Annual budget', minimum=0, maximum=500_000,
        )
        if error:
            errors.append(error)

        target_major = _mapped_value(raw_row, headers_by_key, mapping, 'target_major').strip()
        target_countries = _mapped_value(raw_row, headers_by_key, mapping, 'target_countries').strip()
        if len(target_major) > 160:
            errors.append('Target major must use 160 characters or fewer.')
        if len(target_countries) > 255:
            errors.append('Target countries must use 255 characters or fewer.')

        portfolio_url = str(portfolio_links.get(str(row_number), '') or '').strip()
        if not portfolio_url:
            portfolio_url = _mapped_value(raw_row, headers_by_key, mapping, 'portfolio_google_docs_url').strip()
        if portfolio_url:
            if len(portfolio_url) > PORTFOLIO_URL_MAX_LENGTH:
                errors.append(f'Portfolio link must use {PORTFOLIO_URL_MAX_LENGTH} characters or fewer.')
            elif not _valid_google_docs_url(portfolio_url):
                errors.append('Portfolio must use a valid https://docs.google.com/document/... link.')

        name_key = _identity_key(full_name)
        duplicate_reason = ''
        if email and email in existing_emails:
            duplicate_reason = 'A user with this email already exists.'
        elif name_key and name_key in existing_names:
            duplicate_reason = 'A student with this full name already exists in the school.'
        elif email and email in seen_emails:
            duplicate_reason = 'This email is repeated in the import file.'
        elif name_key and name_key in seen_names:
            duplicate_reason = 'This full name is repeated in the import file.'
        if email:
            seen_emails.add(email)
        if name_key:
            seen_names.add(name_key)

        status = 'error' if errors or global_errors else 'duplicate' if duplicate_reason else 'ready'
        if duplicate_reason:
            warnings.append(duplicate_reason)
        rows.append({
            'row_number': row_number,
            'status': status,
            'errors': errors,
            'warnings': warnings,
            'values': {
                'full_name': full_name,
                'grade': grade or grade_raw,
                'email': email,
                'phone': phone,
                'parent_contact': parent_contact,
                'gpa': str(gpa) if gpa is not None else '',
                'ielts_score': str(ielts_score) if ielts_score is not None else '',
                'sat_score': sat_score if sat_score is not None else '',
                'target_major': target_major,
                'target_countries': target_countries,
                'budget_usd': budget_usd if budget_usd is not None else '',
                'scholarship_needed': _boolean_value(
                    _mapped_value(raw_row, headers_by_key, mapping, 'scholarship_needed'), True,
                ),
                'is_active': _boolean_value(
                    _mapped_value(raw_row, headers_by_key, mapping, 'status'), True,
                ),
                'portfolio_google_docs_url': portfolio_url,
            },
        })

    mapped_columns = {value for value in mapping.values() if value}
    summary = {
        'total': len(rows),
        'ready': sum(row['status'] == 'ready' for row in rows),
        'duplicate': sum(row['status'] == 'duplicate' for row in rows),
        'error': sum(row['status'] == 'error' for row in rows),
    }
    return {
        'sheet_name': table['sheet_name'],
        'headers': [{key: value for key, value in header.items() if key != 'index'} for header in table['headers']],
        'fields': IMPORT_FIELDS,
        'mapping': mapping,
        'suggested_mapping': inferred,
        'ignored_columns': [header['label'] for header in table['headers'] if header['key'] not in mapped_columns],
        'global_errors': global_errors,
        'summary': summary,
        'rows': rows,
    }
